"""Reconstrói e valida o dataset de segurança, exportando uma auditoria por linha.

Uso: python -m scripts.validate_security_dataset
O Excel e os CSVs existentes são somente lidos. Divergências interrompem a execução.
"""
from pathlib import Path
import hashlib
import json

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXPECTED_MD5 = '7213a3ee515a713f4eee2a6948f1756e'
EXPECTED_CLASSES = {
    'Baiting', 'Malware', 'NOT-Malicious General Class',
    'Phishing', 'Pretexting', 'Scareware',
}


def sha(text: str) -> str:
    return hashlib.sha256(text.encode('utf-8')).hexdigest()


def validate_dataset(root: Path = ROOT, export: bool = True) -> dict:
    raw_path = root / 'data/raw/security/phishing_nlp_dataset.xlsx'
    original_bytes = raw_path.read_bytes()
    assert hashlib.md5(original_bytes).hexdigest() == EXPECTED_MD5, 'Excel diferente da fonte conferida.'
    workbook = load_workbook(raw_path, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ['in']
        sheet = workbook['in']
        assert (sheet.max_row, sheet.max_column) == (625, 2)
        excel_rows = list(sheet.iter_rows(values_only=True))
        assert excel_rows[0] == ('Corpus', 'Labels')
        assert not any(c.data_type == 'f' for row in sheet for c in row), 'Fórmulas inesperadas.'
    finally:
        workbook.close()

    raw = pd.read_excel(raw_path, sheet_name='in')
    assert raw.shape == (624, 2) and list(raw.columns) == ['Corpus', 'Labels']
    assert raw['Corpus'].notna().all()
    records, provenance = [], []
    for index, row in raw.iterrows():
        corpus = row['Corpus']
        fragment = row['Labels']
        assert isinstance(corpus, str)
        # Verificação independente da leitura pandas contra as células openpyxl.
        assert corpus == excel_rows[index + 1][0]
        split = pd.notna(fragment)
        if split:
            assert isinstance(fragment, str) and fragment == excel_rows[index + 1][1]
            combined = corpus.rstrip() + ' ' + fragment.lstrip()
        else:
            assert excel_rows[index + 1][1] is None
            combined = corpus
        assert '\t' in combined, f'Rótulo sem delimitador na linha Excel {index + 2}'
        message_part, category_part = combined.rsplit('\t', 1)
        text, category = message_part.strip(), category_part.strip()
        assert text and category in EXPECTED_CLASSES
        records.append({'text': text, 'category': category})
        provenance.append({
            'sheet': 'in', 'excel_row': index + 2,
            'reconstruction': 'join_Corpus_Labels' if split else 'split_Corpus_last_tab',
            'category': category, 'text_sha256': sha(text),
            'normalized_text_sha256': sha(text.strip().lower()),
            'corpus_sha256': sha(corpus),
            'labels_cell_sha256': sha(fragment) if split else '',
            'outer_text_whitespace_removed': text != message_part,
        })

    reconstructed = pd.DataFrame(records)
    audit = pd.DataFrame(provenance)
    keys = reconstructed['text'].str.strip().str.lower()
    labelled = reconstructed.assign(key=keys)
    labels_per_key = labelled.groupby('key')['category'].nunique()
    conflict_keys = set(labels_per_key[labels_per_key.gt(1)].index)
    conflict_mask = keys.isin(conflict_keys)
    first_kept = {}
    clean_indices = []
    decisions, references, group_labels = [], [], []
    labels_by_key = labelled.groupby('key')['category'].agg(lambda v: '|'.join(sorted(set(v))))
    for index, row in labelled.iterrows():
        key = row['key']
        group_labels.append(labels_by_key[key])
        if key in conflict_keys:
            decisions.append('excluded_label_conflict')
            references.append(None)
        elif key in first_kept:
            decisions.append('excluded_duplicate')
            references.append(first_kept[key] + 2)
        else:
            first_kept[key] = index
            clean_indices.append(index)
            decisions.append('kept')
            references.append(index + 2)
    audit['decision'] = decisions
    audit['reference_excel_row'] = pd.array(references, dtype='Int64')
    audit['labels_in_normalized_group'] = group_labels
    index_to_csv_row = {index: position + 2 for position, index in enumerate(clean_indices)}
    audit['processed_csv_row'] = pd.array([index_to_csv_row.get(i) for i in raw.index], dtype='Int64')
    cleaned = reconstructed.loc[clean_indices].reset_index(drop=True)

    # Reprodução independente da regra do notebook do TP1.
    expected = labelled.loc[~conflict_mask].drop_duplicates(['key', 'category'], keep='first')
    pd.testing.assert_frame_equal(cleaned, expected[['text', 'category']].reset_index(drop=True))
    assert reconstructed.shape == (624, 2) and cleaned.shape == (603, 2)
    assert set(cleaned.category) == EXPECTED_CLASSES
    assert int(conflict_mask.sum()) == 15 and len(conflict_keys) == 6
    assert decisions.count('excluded_duplicate') == 6
    assert cleaned.notna().all().all() and cleaned.text.str.strip().ne('').all()
    assert not cleaned.text.str.strip().str.lower().duplicated().any()
    assert len(audit) == len(raw) and audit.excel_row.is_unique
    assert set(audit.excel_row) == set(range(2, 626))
    for source_index, clean_index in zip(clean_indices, range(len(cleaned))):
        assert cleaned.loc[clean_index, 'text'] == reconstructed.loc[source_index, 'text']
        assert cleaned.loc[clean_index, 'category'] == reconstructed.loc[source_index, 'category']

    comparisons = {}
    for relative, expected_frame in [
        ('data/interim/security/phishing_nlp_dataset.csv', reconstructed),
        ('data/processed/security/phishing_nlp_dataset.csv', cleaned),
    ]:
        existing = pd.read_csv(root / relative)
        pd.testing.assert_frame_equal(existing, expected_frame)
        comparisons[relative] = {'equal_values_and_order': True,
                                 'sha256': hashlib.sha256((root / relative).read_bytes()).hexdigest()}

    excluded = audit.loc[audit.decision.ne('kept')].copy()
    counts = pd.DataFrame({
        'before': reconstructed.category.value_counts(),
        'excluded_conflict': audit.loc[audit.decision.eq('excluded_label_conflict'), 'category'].value_counts(),
        'excluded_duplicate': audit.loc[audit.decision.eq('excluded_duplicate'), 'category'].value_counts(),
        'after': cleaned.category.value_counts(),
    }).fillna(0).astype(int).sort_index()
    assert (counts.before - counts.excluded_conflict - counts.excluded_duplicate == counts.after).all()
    counts['percent_after'] = (100 * counts.after / len(cleaned)).round(2)
    summary = {
        'source': 'https://zenodo.org/records/15235123',
        'raw_md5': hashlib.md5(original_bytes).hexdigest(),
        'raw_sha256': hashlib.sha256(original_bytes).hexdigest(),
        'raw_shape': list(raw.shape), 'raw_dtypes': raw.dtypes.astype(str).to_dict(),
        'raw_missing': {k: int(v) for k, v in raw.isna().sum().items()},
        'split_message_excel_rows': audit.loc[audit.reconstruction.eq('join_Corpus_Labels'), 'excel_row'].tolist(),
        'reconstructed_shape': list(reconstructed.shape),
        'reconstructed_dtypes': reconstructed.dtypes.astype(str).to_dict(),
        'reconstructed_missing': {k: int(v) for k, v in reconstructed.isna().sum().items()},
        'exact_duplicate_excess': int(reconstructed.duplicated().sum()),
        'rows_in_exact_duplicates': int(reconstructed.duplicated(keep=False).sum()),
        'rows_in_normalized_duplicates': int(keys.duplicated(keep=False).sum()),
        'conflicting_normalized_groups': len(conflict_keys),
        'excluded_conflict_rows': int(conflict_mask.sum()),
        'excluded_duplicate_rows': decisions.count('excluded_duplicate'),
        'clean_shape': list(cleaned.shape), 'classes': sorted(EXPECTED_CLASSES),
        'clean_dtypes': cleaned.dtypes.astype(str).to_dict(),
        'clean_missing': {k: int(v) for k, v in cleaned.isna().sum().items()},
        'clean_normalized_duplicates': int(cleaned.text.str.strip().str.lower().duplicated().sum()),
        'retained_texts_equal_to_reconstructed': True,
        'text_preservation_scope': 'Preserva texto reconstruído; reconstrução aplica strip e junta fragmentos com um espaço.',
        'existing_csv_comparisons': comparisons,
        'pandas_version': pd.__version__, 'checks_passed': True,
    }
    assert raw_path.read_bytes() == original_bytes
    if export:
        output = root / 'reports/TP2/security_validation'
        output.mkdir(parents=True, exist_ok=True)
        audit.to_csv(output / 'row_audit.csv', index=False)
        excluded.to_csv(output / 'excluded_rows.csv', index=False)
        counts.to_csv(output / 'class_counts.csv', index_label='category')
        reconstructed.describe(include='all').to_csv(output / 'describe_reconstructed.csv', index_label='statistic')
        cleaned.describe(include='all').to_csv(output / 'describe_clean.csv', index_label='statistic')
        (output / 'summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return {'raw': raw, 'reconstructed': reconstructed, 'cleaned': cleaned,
            'audit': audit, 'excluded': excluded, 'counts': counts, 'summary': summary}


def main():
    result = validate_dataset()
    print('Validação concluída: 624 - 15 conflitos - 6 duplicatas = 603 registros; 6 classes.')
    print(result['counts'].to_string())
    print('Excel preservado; CSVs existentes idênticos em valores e ordem; auditoria exportada.')


if __name__ == '__main__':
    main()
